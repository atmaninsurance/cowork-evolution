import struct, os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION

def png_size(path):
    with open(path,'rb') as f:
        f.read(16); w,h = struct.unpack('>II', f.read(8))
    return w,h

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)

def set_portrait(s):
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.top_margin = s.bottom_margin = Inches(0.7)
    s.left_margin = s.right_margin = Inches(0.8)

def set_landscape(s):
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Inches(11), Inches(8.5)
    s.top_margin = s.bottom_margin = Inches(0.5)
    s.left_margin = s.right_margin = Inches(0.5)

set_portrait(doc.sections[0])

def heading(text, size=16):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size)

def body(text, size=10.5, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.font.size = Pt(size); r.italic = italic

def add_map(png, usable_w, usable_h):
    w,h = png_size(png); ar = w/h
    dw = usable_w; dh = dw/ar
    if dh > usable_h: dh = usable_h; dw = dh*ar
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(png, width=Inches(dw))

# Page 1 (portrait): title + arrivals
heading("The Front Door — Designed End-State", 22)
body("As designed through 2026-08-08.  Governing records: DEC-260114 (+ Amendments 1–2), "
     "DEC-260119 (+ Amendment 1), the Verifier design of record, and the ratified "
     "Reviewer/Housekeeper split (OI-000037 r4).", italic=True)
body("The design in one sentence: code proves everything enumerable; exactly two short "
     "single-question model checks guard authorization and fidelity; one lane — the Designer — "
     "holds the pen for everything that requires judgment; every judging pass reads fingerprinted "
     "copies, never the live records; and every path either routes, disposes within a named bound, "
     "or waits on David's board forever — nothing expires, nothing is dropped.", size=11)
heading("1 · What arrives, and where it goes", 14)
tbl = doc.add_table(rows=1, cols=3); tbl.style = 'Light Grid Accent 1'
hdr = tbl.rows[0].cells
for i,t in enumerate(["Arrival","Recognized by","Enters"]): hdr[i].paragraphs[0].add_run(t).bold = True
for r in [("Fresh request","anything not matching the rows below","the intake flow (map 2)"),
          ("David's decision","origin: david-decision naming an open item","decision re-entry (map 5)"),
          ("Designer return","the Designer lane's stamped filing for an item it holds","the second pass (map 4)")]:
    c = tbl.add_row().cells
    for i,t in enumerate(r): c[i].paragraphs[0].add_run(t).font.size = Pt(10)
doc.add_paragraph()
add_map("mmd/m1-arrivals.png", 6.9, 4.2)


# --- Handoff mechanics page (inserted 2026-08-08, David-directed)
sm = doc.add_section(WD_SECTION.NEW_PAGE); set_portrait(sm)
heading("1.1 · Handoff mechanics — how work actually moves", 15)
for para in [
 "A piece of work IS a document, and custody IS the folder it rests in. A handoff is an atomic "
 "move from one folder to another; the receiving lane notices because its folder is watched, so "
 "the move itself is the wake-up call. Two standing rules protect this: compose outside the "
 "folder, then one move (arrival is a single atomic event — no partially-written file is ever "
 "visible), and no lane writes into a folder another lane watches (a folder with two writers has "
 "ambiguous custody).",
 "Item state is the exception — it changes in place. The item's spine (orchestrator/items/OI-N.md) "
 "never moves during its life; status transitions are edits to it, and only the orchestrator holds "
 "that pen. Sidecars (intake, responses, decisions, escalation packages) accumulate beside it, "
 "each made read-only on arrival. The whole set moves exactly once, to the archive, at the end.",
 "The handoff to David is the other exception. Escalation moves nothing: a package is written in "
 "place, the item enters a waiting state, and he is notified (count-only). His return follows the "
 "normal rule — a decision document moved into the inbox.",
 "Terminal handoffs leave the tree. Delivery writes results to their deliver_to destination (a "
 "project folder, the daily log), and the authorization echo appends rulings into the estate "
 "record. Everything between intake and delivery stays inside Agent_Workflow — which is why the "
 "pipeline's entire state survives crashes and is inspectable with a file listing: the filesystem "
 "is the queue.",
]:
    body(para, size=10.5)
doc.add_paragraph()
heading("Map legend — steps are annotated only where documents change; a check that merely reads carries no annotation", 11)
t3 = doc.add_table(rows=1, cols=2); t3.style = 'Light Grid Accent 1'
h3 = t3.rows[0].cells
h3[0].paragraphs[0].add_run("Symbol").bold = True
h3[1].paragraphs[0].add_run("Meaning").bold = True
for a,b in [("＋","document created"),("✎","document modified in place"),("⇒","document moved — source folder → destination folder")]:
    c = t3.add_row().cells
    c[0].paragraphs[0].add_run(a).font.size = Pt(11)
    c[1].paragraphs[0].add_run(b).font.size = Pt(10)
doc.add_paragraph()
body("Folder roster: orchestrator/inbox/ (the front door) · orchestrator/items/ (spines + read-only "
     "sidecars) · orchestrator/archive/ (closed items) · designer/inbox → processing → done (the "
     "Designer's lane) · designer/designs/ (design records) · code/logs/ (composition staging for "
     "tasks) · code/inbox → processing → outbox (transit) → reviewed/ (the executor's lane) · "
     "code/artifacts/ (deliverables) · code/dead-letter/ · _meta/grants/ (capability grants) · "
     "per-call staging (judges' fingerprinted copies, removed after each call) · exports: deliver_to "
     "destinations + The_Estate/.", size=10)

maps = [
 ("2 · A fresh request",
  "Four code checks, then one model question — the authorization check is the only general-purpose "
  "judgment at the front door; everything else is provable by code.",
  "mmd/m2-fresh-request.png", "portrait"),
 ("3 · The Designer lane — where judgment lives",
  "One lane holds the pen: viability, alternatives, the drafted task, executor and model "
  "(raise-only to supervised), variation stated, claim-side questions answered in-lane. Bounded "
  "three ways — round cap (3), ~24 h age-out, and the progress judge — every exit lands on David's board.",
  "mmd/m3-designer-lane.png", "portrait"),
 ("4 · The second pass — a filing that carries a design",
  "Code proves the filing is what the Designer verified; one model question checks it serves the "
  "intent. Auto-route without pinging David requires: consultation origin + authorization yes + "
  "clean code checks + a designed block + the Designer's stated variation: none.",
  "mmd/m4-second-pass.png", "portrait"),
 ("5 · Decision re-entry — David's ruling comes back",
  "The ruling is quarantined immutably and echoed into the estate record before anything is enacted. "
  "A rewrite passes the second pass (code + fidelity question) — no full re-judgment.",
  "mmd/m5-decision-reentry.png", "portrait"),
 ("6 · The executor lane — routing to closure",
  "Review is in the path, not beside it: nothing reaches completion unreviewed. The Reviewer wakes "
  "on events; the Housekeeper wakes on time — events catch things that happen, timers catch things "
  "that fail to happen.",
  "mmd/m6-executor-lane.png", "portrait"),
]
cur = "portrait"
for title, lead, png, orient in maps:
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    (set_landscape if orient == "landscape" else set_portrait)(s)
    cur = orient
    heading(title, 15)
    body(lead, size=10)
    if orient == "landscape": add_map(png, 10.0, 6.3)
    else: add_map(png, 6.9, 7.6)

s = doc.add_section(WD_SECTION.NEW_PAGE); set_portrait(s)
heading("7 · How the judges read — copies, not originals", 15)
body("Every judging pass receives its record inputs as fingerprinted copies in a per-call staging "
     "directory — the spine content, the quarantined sidecars, the resolved governance anchors — and "
     "holds no path into the live estate, orchestrator, decisions, or transcript trees. The "
     "fingerprints are recorded with the verdict, so what a judgment rested on is reproducible.")
body("Three deliberate live reads, and only these: the byte-identity check hashes the REAL filing "
     "that would route (a copy would prove the copy); the fidelity check reads the LIVE design record "
     "in the Designer's own tree; the Designer keeps live-tree code access and write access to its own "
     "lane. The copier fails closed — an unreadable input aborts the launch loudly rather than letting "
     "a pass judge on partial evidence. The write guard watches the staging root; a pass that cannot "
     "reach the records needs no broad fence, and a concurrent writer anywhere in the estate is "
     "invisible to it.")
doc.add_paragraph()
heading("8 · Who decides what", 15)
t2 = doc.add_table(rows=1, cols=2); t2.style = 'Light Grid Accent 1'
h2 = t2.rows[0].cells
h2[0].paragraphs[0].add_run("Decision").bold = True
h2[1].paragraphs[0].add_run("Owner").bold = True
for a,b in [
 ("Is this filing safe and well-formed to look at?","code (screen, red-line, well-formedness)"),
 ("Does the record authorize this, at this capability and scope?","authorization check (one model question)"),
 ("Is this worth doing? Can it be done? How? Alternatives?","Designer"),
 ("What questions block this, and are any David's?","Designer (only David-class questions leave the lane)"),
 ("Which executor, which model, attended or not?","Designer (may raise to supervised, never lower)"),
 ("Does it deviate from what the record contemplates?","Designer states it; a significant variation blocks auto-route"),
 ("Is what would run what was authorized?","code (byte-identity) + fidelity check (one model question)"),
 ("Is this alert a false alarm?","Designer recommends; orchestrator enacts — machine-raised alerts only"),
 ("Did the finished work pass muster?","Reviewer (approve / return for work / return for redesign)"),
 ("Is the lane healthy?","Housekeeper (timed; watches for what failed to happen)"),
 ("Everything escalated","David — rules in-chat; the decision document is the enactment"),
]:
    c = t2.add_row().cells
    c[0].paragraphs[0].add_run(a).font.size = Pt(10)
    c[1].paragraphs[0].add_run(b).font.size = Pt(10)
doc.add_page_break()
heading("9 · Invariants — true on every path", 15)
for inv in [
 "Quarantine-first. Every consumed arrival becomes an immutable sidecar before anything acts on it; only the orchestrator writes the item's spine.",
 "Screen at every promotion. Re-screened at every hop toward execution, whoever authored it.",
 "Author ≠ checker. The session that writes a verdict never held the pen, and vice versa.",
 "Judges cannot touch the evidence. Structurally — they read copies.",
 "Nothing routes on a guess. An unresolvable cited record is a loud failure, never a fuzzy match.",
 "Review precedes completion. Nothing is delivered unreviewed.",
 "Authorization echo. A consumed ruling is written into the estate record before enactment.",
 "Fail toward David, never toward silence. Every doubt, cap, age-out, and refusal lands on the board and waits; escalations never expire; his answers arrive only through decision documents.",
 "Every loop is bounded. Adjudication retries, Designer rounds, review return-cycles, worker attempts, age-outs — a loop with no bound is a defect by definition.",
]:
    p = doc.add_paragraph(style='List Bullet'); r = p.add_run(inv); r.font.size = Pt(10.5)

out = os.path.expanduser("~/Documents/Projects/cowork-evolution/Design/front-door-endstate-20260808.docx")
doc.save(out)
print("saved:", out, os.path.getsize(out), "bytes")
